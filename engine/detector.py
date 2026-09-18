"""
detector.py - Motor de identificación automática de archivos PAI.
Clasifica los archivos independientemente del nombre con el que el municipio los guarde.
"""
import openpyxl
import os
import re

MUNICIPIOS_NOMBRES = [
    "PEREIRA", "APIA", "BALBOA", "BELEN DE UMBRIA", "DOSQUEBRADAS",
    "GUATICA", "LA CELIA", "LA VIRGINIA", "MARSELLA", "MISTRATO",
    "PUEBLO RICO", "QUINCHIA", "SANTA ROSA DE CABAL", "SANTUARIO"
]

def normalizar_texto(texto):
    if not texto:
        return ""
    t = str(texto).upper()
    reemplazos = (
        ("Á", "A"), ("É", "E"), ("Í", "I"), ("Ó", "O"), ("Ú", "U"),
        ("Ü", "U"), ("Ñ", "N")
    )
    for a, b in reemplazos:
        t = t.replace(a, b)
    return re.sub(r'[^A-Z0-9\s]', ' ', t).strip()

def detectar_tipo_archivo(filepath):
    """
    Analiza un archivo Excel y determina su tipo y municipio.
    Retorna dict con:
      - tipo: 'DOSIS_APLICADAS' | 'MOVIMIENTO_BIOLOGICOS' | 'VACUNADOS_EXTRANJEROS' | 'DESCONOCIDO'
      - municipio: nombre del municipio detectado o None
      - hojas: lista de hojas del libro
      - confianza: 'ALTA' | 'MEDIA' | 'BAJA'
      - descripcion: texto descriptivo
    """
    resultado = {
        "tipo": "DESCONOCIDO",
        "municipio": None,
        "hojas": [],
        "confianza": "BAJA",
        "descripcion": ""
    }

    try:
        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
        hojas = wb.sheetnames
        resultado["hojas"] = hojas
        hojas_norm = [normalizar_texto(h) for h in hojas]

        # 1. ¿Es VACUNADOS EXTRANJEROS?
        if any("VENEZOLANOS" in h or "BRASIL" in h or "ECUADOR" in h for h in hojas_norm):
            resultado["tipo"] = "VACUNADOS_EXTRANJEROS"
            resultado["confianza"] = "ALTA"
            resultado["descripcion"] = "Plantilla de Vacunados Extranjeros (Países Fronterizos)"

        # 2. ¿Es DOSIS APLICADAS?
        elif any("PLANTILLA MENSUAL" in h or "CONSOLIDADO REGIMEN" in h or "COBERTURAXMPIOS" in h for h in hojas_norm):
            resultado["tipo"] = "DOSIS_APLICADAS"
            resultado["confianza"] = "ALTA"
            resultado["descripcion"] = "Plantilla Mensual de Dosis Aplicadas PAI"

        # 3. ¿Es MOVIMIENTO DE BIOLÓGICOS (Kardex / MB)?
        else:
            meses_encontrados = sum(1 for h in hojas_norm if any(m in h for m in [
                "ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO",
                "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"
            ]))
            if meses_encontrados >= 3 or any("PERDIDAS" in h or "VALIDADOR" in h for h in hojas_norm):
                resultado["tipo"] = "MOVIMIENTO_BIOLOGICOS"
                resultado["confianza"] = "ALTA"
                resultado["descripcion"] = "Movimiento Mensual de Biológicos e Insumos (Kardex/MB)"

        # Detectar municipio en el nombre del archivo
        nombre_archivo_norm = normalizar_texto(os.path.basename(filepath))
        for m in MUNICIPIOS_NOMBRES:
            m_norm = normalizar_texto(m)
            if m_norm in nombre_archivo_norm:
                resultado["municipio"] = m
                break

        # Si no lo encontró en el nombre, buscar en las primeras filas de la primera hoja
        if not resultado["municipio"] and hojas:
            ws = wb[hojas[0]]
            for row in ws.iter_rows(max_row=10, max_col=10, values_only=True):
                for cell in row:
                    if cell:
                        c_norm = normalizar_texto(cell)
                        for m in MUNICIPIOS_NOMBRES:
                            if normalizar_texto(m) in c_norm:
                                resultado["municipio"] = m
                                break
                    if resultado["municipio"]:
                        break
                if resultado["municipio"]:
                    break

        wb.close()

    except Exception as e:
        resultado["descripcion"] = f"Error al abrir Excel: {str(e)}"

    return resultado
