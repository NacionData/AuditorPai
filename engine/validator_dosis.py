"""
validator_dosis.py - Validador con Recalculo Manual Independiente para Dosis Aplicadas PAI.
Audita la consistencia recalculando fila por fila para evitar que formulas alteradas
o celdas sobreescritas oculten descuadros.
"""
import openpyxl
import re

MESES_ORDEN = [
    "ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO",
    "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"
]

def normalizar(texto):
    if not texto:
        return ""
    t = str(texto).upper().strip()
    for a, b in [("Á", "A"), ("É", "E"), ("Í", "I"), ("Ó", "O"), ("Ú", "U"), ("Ñ", "N")]:
        t = t.replace(a, b)
    return t

def safe_num(val):
    if val is None or val == "":
        return 0
    try:
        if isinstance(val, (int, float)):
            return int(round(val))
        v_str = str(val).strip().replace(",", ".")
        return int(round(float(v_str)))
    except (ValueError, TypeError):
        return None

def validar_dosis(filepath, mes_evaluar="AGOSTO", municipio_nombre=None):
    """
    Audita el archivo de Dosis Aplicadas para el mes especificado.
    RECALCULA manualmente cada sumatoria (Género, Régimen, Étnicos) en lugar
    de confiar en las celdas de fórmula, detectando adulteraciones de fórmulas.
    """
    mes_evaluar = normalizar(mes_evaluar)
    resultado = {
        "valido": True,
        "tipo": "DOSIS_APLICADAS",
        "mes": mes_evaluar,
        "municipio": municipio_nombre or "DESCONOCIDO",
        "total_dosis_mes": 0,
        "total_vacunas_con_datos": 0,
        "errores": [],
        "advertencias": [],
        "resumen_coherencia": {
            "columnas_evaluadas": 0,
            "columnas_con_errores": 0,
            "coincidencia_genero_regimen": True,
            "coincidencia_genero_etnico": True,
            "formulas_adulteradas": 0
        }
    }

    try:
        wb = openpyxl.load_workbook(filepath, data_only=True, read_only=True)
        if "1_PLANTILLA_MENSUAL" not in wb.sheetnames:
            resultado["valido"] = False
            resultado["errores"].append({
                "tipo": "HOJA_FALTANTE",
                "mensaje": "El archivo no contiene la hoja obligatoria '1_PLANTILLA_MENSUAL'."
            })
            wb.close()
            return resultado

        ws = wb["1_PLANTILLA_MENSUAL"]

        # Carga en memoria de las 541 columnas oficiales MinSalud 2026
        max_r = 260
        max_c = 541
        grid = []
        for row in ws.iter_rows(min_row=1, max_row=max_r, min_col=1, max_col=max_c, values_only=True):
            grid.append(row)
        wb.close()

        dosis_por_columna = {}

        def get_val(r, c):
            if 1 <= r <= len(grid):
                row = grid[r - 1]
                if 1 <= c <= len(row):
                    return row[c - 1]
            return None

        # 1. Leer encabezados de vacunas
        col_nombres = {}
        col_familias = {}
        curr_vacuna = ""
        for c_idx in range(5, max_c + 1):
            v6 = get_val(6, c_idx)
            v7 = get_val(7, c_idx)
            v8 = get_val(8, c_idx)
            if v6:
                curr_vacuna = str(v6).strip()
            col_familias[c_idx] = curr_vacuna
            detalle = " - ".join([str(p).strip() for p in [v7, v8] if p and str(p).strip()])
            nombre_completo = f"{curr_vacuna} ({detalle})" if detalle else (curr_vacuna or f"Columna {c_idx}")
            col_nombres[c_idx] = re.sub(r'\s+', ' ', nombre_completo).strip()

        # 1.1 Verificación de Vigencia y Estructura Oficial 2026
        ano_detectado = None
        for r_check in range(1, 6):
            for c_check in range(1, 10):
                val_check = get_val(r_check, c_check)
                if val_check:
                    val_str = str(val_check).upper()
                    match_ano = re.search(r'\b(202[0-5])\b', val_str)
                    if match_ano:
                        ano_detectado = int(match_ano.group(1))
                        break
            if ano_detectado:
                break

        tiene_hexavalente = any("HEXAVALENTE" in str(col_nombres.get(c, "")).upper() for c in col_nombres)

        if (ano_detectado and ano_detectado < 2026) or (not tiene_hexavalente and len(col_nombres) > 15):
            resultado["valido"] = False
            resultado["errores"].append({
                "tipo": "PLANTILLA_DESACTUALIZADA",
                "mensaje": f"PLANTILLA DESACTUALIZADA: El archivo subido corresponde a una vigencia anterior ({ano_detectado or 'previa a 2026'}). MinSalud modificó la estructura oficial para 2026 incorporando el nuevo esquema (Hexavalente, VSR) y columnas reestructuradas. Por favor descarga y diligencia la Plantilla Oficial 2026 en blanco desde el botón superior."
            })
            return resultado

        # 2. Localizar el bloque de filas del mes
        mes_idx = MESES_ORDEN.index(mes_evaluar) if mes_evaluar in MESES_ORDEN else 7
        fila_teorica = 9 + (mes_idx * 19)

        start_row = None
        val_c3 = get_val(fila_teorica, 3)
        if val_c3 and mes_evaluar in normalizar(val_c3):
            start_row = fila_teorica
        else:
            for r in range(8, len(grid) + 1):
                c3 = get_val(r, 3)
                c2 = get_val(r, 2)
                c4 = get_val(r, 4)
                if c3 and mes_evaluar in normalizar(c3):
                    start_row = r
                    break
                elif c2 and str(c2).strip() == str(mes_idx + 1) and c4 and "FEMENINO" in normalizar(c4):
                    start_row = r
                    break

        if not start_row:
            start_row = fila_teorica

        total_dosis_acumulado = 0
        columnas_evaluadas = 0
        columnas_con_error = 0
        formulas_alteradas = 0
        vacunas_agrupadas = {}

        for col_idx in range(5, max_c + 1):
            # Omitir columnas auxiliares de metadatos o resúmenes de plantilla
            v7_raw = str(get_val(7, col_idx) or "").upper()
            v8_raw = str(get_val(8, col_idx) or "").upper()
            if "CODEPCOMURE" in v7_raw or "TOTAL DOSIS" in v8_raw or "DOSIS APLICADAS POR BIOLOGICOS" in v7_raw:
                continue
            if col_idx in [474, 475, 520, 521, 524]:
                continue

            col_letter = openpyxl.utils.get_column_letter(col_idx)
            vacuna_nombre = col_nombres.get(col_idx, f"Columna {col_letter}")

            # --- FILAS DE GÉNERO ---
            val_fem = safe_num(get_val(start_row, col_idx))
            val_masc = safe_num(get_val(start_row + 1, col_idx))
            val_otro = safe_num(get_val(start_row + 2, col_idx))
            celda_tot_gen = safe_num(get_val(start_row + 3, col_idx))

            # --- FILAS DE RÉGIMEN ---
            val_contr = safe_num(get_val(start_row + 4, col_idx))
            val_subsid = safe_num(get_val(start_row + 5, col_idx))
            val_pobre = safe_num(get_val(start_row + 6, col_idx))
            val_especial = safe_num(get_val(start_row + 7, col_idx))
            celda_tot_reg = safe_num(get_val(start_row + 8, col_idx))

            # --- FILAS DE GRUPOS ÉTNICOS ---
            val_indigena = safe_num(get_val(start_row + 9, col_idx))
            val_rom = safe_num(get_val(start_row + 10, col_idx))
            val_raizal = safe_num(get_val(start_row + 11, col_idx))
            val_palenquero = safe_num(get_val(start_row + 12, col_idx))
            val_afro = safe_num(get_val(start_row + 13, col_idx))
            val_sin_etnia = safe_num(get_val(start_row + 14, col_idx))
            celda_tot_etn = safe_num(get_val(start_row + 15, col_idx))

            todos_valores = [
                val_fem, val_masc, val_otro, celda_tot_gen,
                val_contr, val_subsid, val_pobre, val_especial, celda_tot_reg,
                val_indigena, val_rom, val_raizal, val_palenquero, val_afro, val_sin_etnia, celda_tot_etn
            ]

            # Verificar si hay celdas de texto corruptas
            if any(v is None for v in todos_valores):
                resultado["errores"].append({
                    "tipo": "TEXTO_EN_CAMPO_NUMERICO",
                    "columna": col_letter,
                    "vacuna": vacuna_nombre,
                    "mensaje": f"En '{vacuna_nombre}' (Col {col_letter}): Se encontró texto o caracteres inválidos en celdas que deben ser números."
                })
                columnas_con_error += 1
                resultado["valido"] = False
                continue

            v_fem = val_fem or 0
            v_masc = val_masc or 0
            v_otro = val_otro or 0
            c_tot_gen = celda_tot_gen or 0

            v_contr = val_contr or 0
            v_subsid = val_subsid or 0
            v_pobre = val_pobre or 0
            v_especial = val_especial or 0
            c_tot_reg = celda_tot_reg or 0

            v_indigena = val_indigena or 0
            v_rom = val_rom or 0
            v_raizal = val_raizal or 0
            v_palenquero = val_palenquero or 0
            v_afro = val_afro or 0
            v_sin_etnia = val_sin_etnia or 0
            c_tot_etn = celda_tot_etn or 0

            # Si toda la columna está en ceros, saltar
            if all(v == 0 for v in [v_fem, v_masc, v_otro, c_tot_gen, v_contr, v_subsid, v_pobre, v_especial, c_tot_reg, v_indigena, v_rom, v_raizal, v_palenquero, v_afro, v_sin_etnia, c_tot_etn]):
                continue

            columnas_evaluadas += 1
            resultado["total_vacunas_con_datos"] += 1

            # 1. Chequeo de números negativos
            if min([v_fem, v_masc, v_otro, v_contr, v_subsid, v_pobre, v_especial, v_indigena, v_rom, v_raizal, v_palenquero, v_afro, v_sin_etnia]) < 0:
                resultado["errores"].append({
                    "tipo": "VALOR_NEGATIVO",
                    "columna": col_letter,
                    "vacuna": vacuna_nombre,
                    "mensaje": f"En '{vacuna_nombre}' (Col {col_letter}): Se registraron números negativos de dosis."
                })
                resultado["valido"] = False
                columnas_con_error += 1
                continue

            # 2. RECALCULO MANUAL INDEPENDIENTE (Sin depender de las celdas de fórmula)
            real_sum_genero = v_fem + v_masc + v_otro
            real_sum_regimen = v_contr + v_subsid + v_pobre + v_especial
            real_sum_etnicos = v_indigena + v_rom + v_raizal + v_palenquero + v_afro + v_sin_etnia

            total_dosis_acumulado += real_sum_genero
            dosis_por_columna[col_idx] = real_sum_genero
            fam_vacuna = col_familias.get(col_idx, "")
            if fam_vacuna:
                vacunas_agrupadas[fam_vacuna] = vacunas_agrupadas.get(fam_vacuna, 0) + real_sum_genero

            # 3. Detección de Fórmulas Adulteradas / Sobreescritas
            # Si el municipio alteró la celda de total escribiendo un número que no es la suma real:
            if c_tot_gen != real_sum_genero:
                formulas_alteradas += 1
                resultado["errores"].append({
                    "tipo": "FORMULA_GENERO_ALTERADA",
                    "columna": col_letter,
                    "vacuna": vacuna_nombre,
                    "mensaje": f"En '{vacuna_nombre}' (Col {col_letter}): La celda TOTAL GÉNERO contiene {c_tot_gen}, pero la suma real (Fem:{v_fem} + Masc:{v_masc} + Otro:{v_otro}) es {real_sum_genero}. Se detectó fórmula alterada o sobreescrita."
                })
                resultado["valido"] = False
                columnas_con_error += 1

            if c_tot_reg != real_sum_regimen:
                formulas_alteradas += 1
                resultado["errores"].append({
                    "tipo": "FORMULA_REGIMEN_ALTERADA",
                    "columna": col_letter,
                    "vacuna": vacuna_nombre,
                    "mensaje": f"En '{vacuna_nombre}' (Col {col_letter}): La celda TOTAL RÉGIMEN contiene {c_tot_reg}, pero la suma real (Contr:{v_contr} + Sub:{v_subsid} + Pobre:{v_pobre} + Esp:{v_especial}) es {real_sum_regimen}. Se detectó fórmula alterada o sobreescrita."
                })
                resultado["valido"] = False
                columnas_con_error += 1

            if c_tot_etn != real_sum_etnicos:
                formulas_alteradas += 1
                resultado["errores"].append({
                    "tipo": "FORMULA_ETNICO_ALTERADA",
                    "columna": col_letter,
                    "vacuna": vacuna_nombre,
                    "mensaje": f"En '{vacuna_nombre}' (Col {col_letter}): La celda TOTAL ÉTNICOS contiene {c_tot_etn}, pero la suma real de grupos étnicos es {real_sum_etnicos}. Se detectó fórmula alterada o sobreescrita."
                })
                resultado["valido"] = False
                columnas_con_error += 1

            # 4. REGLA DE ORO PAI SOBRE VALORES REALES:
            # sum_genero == sum_regimen == sum_etnicos
            if real_sum_genero != real_sum_regimen:
                dif = abs(real_sum_genero - real_sum_regimen)
                resultado["errores"].append({
                    "tipo": "DESCUADRE_GENERO_VS_REGIMEN",
                    "columna": col_letter,
                    "vacuna": vacuna_nombre,
                    "real_genero": real_sum_genero,
                    "real_regimen": real_sum_regimen,
                    "diferencia": dif,
                    "mensaje": f"En '{vacuna_nombre}' (Col {col_letter}): Suma real de Género ({real_sum_genero}) != Suma real de Régimen ({real_sum_regimen}). Descuadre real: {dif} dosis."
                })
                resultado["valido"] = False
                resultado["resumen_coherencia"]["coincidencia_genero_regimen"] = False
                columnas_con_error += 1

            if real_sum_genero != real_sum_etnicos:
                dif = abs(real_sum_genero - real_sum_etnicos)
                resultado["errores"].append({
                    "tipo": "DESCUADRE_GENERO_VS_ETNICO",
                    "columna": col_letter,
                    "vacuna": vacuna_nombre,
                    "real_genero": real_sum_genero,
                    "real_etnico": real_sum_etnicos,
                    "diferencia": dif,
                    "mensaje": f"En '{vacuna_nombre}' (Col {col_letter}): Suma real de Género ({real_sum_genero}) != Suma real de Grupos Étnicos ({real_sum_etnicos}). Descuadre real: {dif} dosis."
                })
                resultado["valido"] = False
                resultado["resumen_coherencia"]["coincidencia_genero_etnico"] = False
                columnas_con_error += 1

        resultado["total_dosis_mes"] = total_dosis_acumulado
        resultado["dosis_agrupadas_vacuna"] = vacunas_agrupadas
        resultado["resumen_coherencia"]["columnas_evaluadas"] = columnas_evaluadas
        resultado["resumen_coherencia"]["columnas_con_errores"] = columnas_con_error
        resultado["resumen_coherencia"]["formulas_adulteradas"] = formulas_alteradas

        # =================================================================
        # 5. INFORME DE ANÁLISIS DE SIMULTANEIDAD DEL ESQUEMA NACIONAL
        # Evaluación pedagógica de las 6 cohortes clave sin bloquear radicación.
        # =================================================================
        def get_d(col):
            return dosis_por_columna.get(col, 0)

        cohortes = [
            {
                "nombre": "Cohorte de 2 Meses",
                "biologicos": [
                    {"nombre": "Pentavalente 1ª + Hexavalente 1ª", "dosis": get_d(102) + get_d(99)},
                    {"nombre": "Antipolio VIP 1ª", "dosis": get_d(67)},
                    {"nombre": "Rotavirus 1ª", "dosis": get_d(157)},
                    {"nombre": "Neumococo 1ª", "dosis": get_d(159)}
                ]
            },
            {
                "nombre": "Cohorte de 4 Meses",
                "biologicos": [
                    {"nombre": "Pentavalente 2ª + Hexavalente 2ª", "dosis": get_d(103) + get_d(100)},
                    {"nombre": "Antipolio VIP 2ª", "dosis": get_d(68)},
                    {"nombre": "Rotavirus 2ª", "dosis": get_d(158)},
                    {"nombre": "Neumococo 2ª", "dosis": get_d(160)}
                ]
            },
            {
                "nombre": "Cohorte de 6 Meses",
                "biologicos": [
                    {"nombre": "Pentavalente 3ª + Hexavalente 3ª", "dosis": get_d(104) + get_d(101)},
                    {"nombre": "Antipolio VIP 3ª", "dosis": get_d(69)},
                    {"nombre": "Influenza Pediátrica 1ª", "dosis": get_d(354)}
                ]
            },
            {
                "nombre": "Cohorte de 12 Meses (1 Año)",
                "biologicos": [
                    {"nombre": "Triple Viral (SRP) 1ª", "dosis": get_d(165)},
                    {"nombre": "Varicela 1ª", "dosis": get_d(264)},
                    {"nombre": "Neumococo Refuerzo", "dosis": get_d(161)},
                    {"nombre": "Hepatitis A Única", "dosis": get_d(251)}
                ]
            },
            {
                "nombre": "Cohorte de 18 Meses (1 Año y Medio)",
                "biologicos": [
                    {"nombre": "Pentavalente 1er Refuerzo", "dosis": get_d(105)},
                    {"nombre": "Antipolio VIP 1er Refuerzo", "dosis": get_d(85)},
                    {"nombre": "Triple Viral (SRP) 2ª", "dosis": get_d(175)}
                ]
            },
            {
                "nombre": "Cohorte de 5 Años",
                "biologicos": [
                    {"nombre": "DPT 2º Refuerzo", "dosis": get_d(125)},
                    {"nombre": "Antipolio VIP 2º Refuerzo", "dosis": get_d(90)},
                    {"nombre": "Varicela Refuerzo", "dosis": get_d(276)}
                ]
            }
        ]

        informe_simultaneidad = []
        for c_info in cohortes:
            nombre_coh = c_info["nombre"]
            bios = c_info["biologicos"]
            dosis_valores = [b["dosis"] for b in bios]
            
            # Solo analizamos si hay actividad en la cohorte
            if any(d > 0 for d in dosis_valores):
                resumen_bios = ", ".join([f"{b['nombre']} ({b['dosis']})" for b in bios])
                max_d = max(dosis_valores)
                min_d = min(dosis_valores)
                
                if max_d == min_d:
                    estado_sim = "OPTIMO"
                    msg = f"Simultaneidad 100% óptima: {resumen_bios}."
                else:
                    estado_sim = "OPORTUNIDAD_OBSERVADA"
                    dif = max_d - min_d
                    msg = f"Desfase observado: {resumen_bios} (diferencia de {dif} dosis entre biológicos)."
                    # Agregamos como advertencia informativa (NO BLOQUEA RADICACIÓN)
                    resultado["advertencias"].append({
                        "tipo": "SIMULTANEIDAD_INFORMATIVA",
                        "cohorte": nombre_coh,
                        "mensaje": f"[Simultaneidad - {nombre_coh}] {resumen_bios}. Desfase observado de {dif} dosis entre biológicos. Se sugiere verificar oportunidades de vacunación o diferimientos clínicos (Informativo, no bloquea radicación)."
                    })
                
                informe_simultaneidad.append({
                    "cohorte": nombre_coh,
                    "estado": estado_sim,
                    "biologicos": bios,
                    "resumen": msg
                })

        resultado["resumen_coherencia"]["informe_simultaneidad"] = informe_simultaneidad

    except Exception as e:
        resultado["valido"] = False
        resultado["errores"].append({
            "tipo": "ERROR_LECTURA",
            "mensaje": f"Fallo al procesar archivo de Dosis: {str(e)}"
        })

    return resultado
