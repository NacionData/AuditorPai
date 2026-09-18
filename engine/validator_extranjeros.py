"""
validator_extranjeros.py - Validador para la plantilla de Vacunados Extranjeros.
Audita las 6 hojas de países fronterizos y su coherencia de reporte mensual.
"""
import openpyxl
import re

HOJAS_PAISES = [
    "1_BRASIL", "2_ECUADOR", "3_PANAMA", "4_PERU", "5_OTROS", "6_VENEZOLANOS"
]

MESES_ORDEN = [
    "ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO",
    "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"
]

def normalizar(texto):
    if not texto:
        return ""
    t = str(texto).upper().strip()
    for a, b in [("Á", "A"), ("É", "E"), ("Í", "I"), ("Ó", "O"), ("Ú", "U"), ("Ñ", "N")]:
        t = t.replace(a, b)
    return t

def safe_num(val):
    if val is None or val == "":
        return 0
    try:
        if isinstance(val, (int, float)):
            return int(round(val))
        return int(round(float(str(val).strip().replace(",", "."))))
    except (ValueError, TypeError):
        return None

def validar_extranjeros(filepath, mes_evaluar="AGOSTO", municipio_nombre=None):
    mes_evaluar = normalizar(mes_evaluar)
    resultado = {
        "valido": True,
        "tipo": "VACUNADOS_EXTRANJEROS",
        "mes": mes_evaluar,
        "municipio": municipio_nombre or "DESCONOCIDO",
        "total_extranjeros_vacunados": 0,
        "hojas_evaluadas": [],
        "errores": [],
        "advertencias": []
    }

    try:
        wb = openpyxl.load_workbook(filepath, data_only=True, read_only=True)
        hojas_disponibles = [h for h in wb.sheetnames]
        hojas_norm = [normalizar(h) for h in hojas_disponibles]

        # Verificar presencia de hojas de países
        hojas_a_revisar = []
        for p in HOJAS_PAISES:
            p_norm = normalizar(p)
            for idx, h in enumerate(hojas_norm):
                if p_norm in h or p_norm.split("_")[-1] in h:
                    hojas_a_revisar.append(hojas_disponibles[idx])
                    break

        if not hojas_a_revisar:
            resultado["valido"] = False
            resultado["errores"].append({
                "tipo": "HOJAS_PAISES_FALTANTES",
                "mensaje": "El archivo no contiene las hojas requeridas de países fronterizos (Venezuela, Brasil, etc.)."
            })
            wb.close()
            return resultado

        # Evaluar cada hoja de país
        mes_idx = MESES_ORDEN.index(mes_evaluar) if mes_evaluar in MESES_ORDEN else 7
        fila_teorica = 9 + (mes_idx * 19)

        total_dosis_ext = 0

        for h_name in hojas_a_revisar:
            ws = wb[h_name]
            resultado["hojas_evaluadas"].append(h_name)

            # Cargar slice en memoria
            grid = []
            for row in ws.iter_rows(min_row=1, max_row=260, min_col=1, max_col=100, values_only=True):
                grid.append(row)

            def get_val(r, c):
                if 1 <= r <= len(grid):
                    row = grid[r - 1]
                    if 1 <= c <= len(row):
                        return row[c - 1]
                return None

            start_row = fila_teorica
            # Comprobar si fila coincide
            val_c3 = get_val(fila_teorica, 3)
            if not (val_c3 and mes_evaluar in normalizar(val_c3)):
                for r in range(8, len(grid) + 1):
                    c3 = get_val(r, 3)
                    if c3 and mes_evaluar in normalizar(c3):
                        start_row = r
                        break

            row_tot_gen = start_row + 3
            row_tot_reg = start_row + 8
            row_tot_etn = start_row + 15

            for c in range(5, 100):
                tot_gen = safe_num(get_val(row_tot_gen, c)) or 0
                tot_reg = safe_num(get_val(row_tot_reg, c)) or 0
                tot_etn = safe_num(get_val(row_tot_etn, c)) or 0

                if tot_gen == 0 and tot_reg == 0 and tot_etn == 0:
                    continue

                total_dosis_ext += tot_gen

                if tot_gen != tot_reg:
                    col_let = openpyxl.utils.get_column_letter(c)
                    resultado["errores"].append({
                        "tipo": "DESCUADRE_EXTRANJEROS",
                        "hoja": h_name,
                        "columna": col_let,
                        "mensaje": f"En hoja '{h_name}' (Col {col_let}): Total Género ({tot_gen}) != Total Régimen ({tot_reg})."
                    })
                    resultado["valido"] = False

        resultado["total_extranjeros_vacunados"] = total_dosis_ext
        wb.close()

    except Exception as e:
        resultado["valido"] = False
        resultado["errores"].append({
            "tipo": "ERROR_LECTURA",
            "mensaje": f"Fallo al procesar archivo de Extranjeros: {str(e)}"
        })

    return resultado
